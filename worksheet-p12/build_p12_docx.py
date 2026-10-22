"""
Script generator lengkap untuk membuat file docx Worksheet P12:
AMADEUS DHARMA AKBARINDRA - 25523253 - 12. Worksheet_PABW_Pertemuan12 - ID.docx

Mencakup seluruh halaman (Cover/Page 1 sampai Lampiran Git Page 23) dengan format, section, layout,
font family (Carlito, DejaVu Sans Mono), style, size, dan isi terisi lengkap 100% presisi.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=65, bottom=65, left=90, right=90):
    """Set cell margins in twips (20 twips = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    """Set borders for the whole table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_callout_box(doc, text, bold_prefix="", hex_bg="EEF2FA", border_color="003796", text_color="1F2328"):
    """Create a callout box styled like the worksheet boxes."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.77)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=75, bottom=75, left=100, right=100)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Carlito"
        r_pre.font.size = Pt(8.5)
        r_pre.font.bold = True
        if border_color == "003796":
            r_pre.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
        elif border_color == "991B1B":
            r_pre.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
        elif border_color == "16A34A":
            r_pre.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
        else:
            r_pre.font.color.rgb = RGBColor(0x85, 0x4D, 0x0E)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Carlito"
    r_txt.font.size = Pt(8.5)
    r_txt.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)

def make_worksheet_header(doc, code, letter, title):
    """Create the distinctive 3-cell worksheet header bar."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(1.5), Inches(0.55), Inches(4.72)]
    for i, w in enumerate(col_widths):
        table.rows[0].cells[i].width = w
        
    cell0 = table.rows[0].cells[0]
    cell1 = table.rows[0].cells[1]
    cell2 = table.rows[0].cells[2]
    
    set_cell_background(cell0, "003796")
    set_cell_background(cell1, "F6D11B")
    set_cell_background(cell2, "EEF2FA")
    
    for c in [cell0, cell1, cell2]:
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(code)
    r0.font.name = "Carlito"
    r0.font.size = Pt(10.0)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    p1 = cell1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(letter)
    r1.font.name = "Carlito"
    r1.font.size = Pt(10.0)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
    
    p2 = cell2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(title)
    r2.font.name = "Carlito"
    r2.font.size = Pt(9.0)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)

def add_code_block(doc, code_text):
    """Add a shaded code block with monospace font."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.77)
    set_cell_background(cell, "F6F8FA")
    set_cell_margins(cell, top=65, bottom=65, left=90, right=90)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:left w:val="single" w:sz="14" w:space="0" w:color="003796"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    r = p.add_run(code_text)
    r.font.name = "DejaVu Sans Mono"
    r.font.size = Pt(8.0)
    r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)

def add_heading_section(doc, text):
    """Add a subsection title in navy blue."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Carlito"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x37, 0x96)

def add_subtext(doc, text, italic=False):
    """Add a descriptive sentence."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Carlito"
    r.font.size = Pt(8.5)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(0x33, 0x38, 0x42)

def add_git_commit_box(doc, commit_letter, commit_msg, note=""):
    """Add the standardized Git commit instruction box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.77)
    set_cell_background(cell, "F6F8FA")
    set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="003796"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    
    r_hdr = p.add_run(f"▌ SIMPAN KE GIT — SETELAH LEMBAR {commit_letter}\n")
    r_hdr.font.name = "Carlito"
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
    
    r_cmd = p.add_run(f"git add .\ngit commit -m \"{commit_msg}\"\ngit push\n\n# pesan commit yang disarankan:\n# {commit_msg}")
    r_cmd.font.name = "DejaVu Sans Mono"
    r_cmd.font.size = Pt(8.0)
    r_cmd.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
    
    default_note = f"Lembar {commit_letter} sudah selesai? Simpan hasilnya sekarang, jangan menunggu kelas berakhir. Simpanan ini tidak memerlukan langkah tambahan apa pun — hanya memastikan pekerjaan Anda tidak bergantung pada satu berkas di satu komputer."
    add_subtext(doc, note if note else default_note, italic=True)

def add_check_verification(doc, letter, text):
    """Add the check verification section."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"{letter} — simpan perubahan\n")
    r_title.font.name = "Carlito"
    r_title.font.size = Pt(9.0)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
    
    r_desc = p.add_run(f"Cara memastikan: {text}")
    r_desc.font.name = "Carlito"
    r_desc.font.size = Pt(8.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(0x33, 0x38, 0x42)

def style_standard_table(table, col_widths, headers, data_rows, header_bg="003796", alt_bg="F8FAFC"):
    """Format a standard table with navy header and neat cells."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D0D7DE", sz="4")
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=55, bottom=55, left=75, right=75)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.name = "Carlito"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    needed_rows = 1 + len(data_rows)
    while len(table.rows) < needed_rows:
        table.add_row()
    while len(table.rows) > needed_rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)

    # Data rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=75, right=75)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.name = "Carlito"
            r.font.size = Pt(8.0)
            r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)


def generate_p12_docx(output_path):
    print("Mulai membuat dokumen P12 terisi penuh...")
    doc = Document()
    
    # 1. Page Setup (A4, precise margins)
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.67)
    sec.bottom_margin = Inches(0.63)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    
    # ==========================================
    # HALAMAN 1: OVERVIEW WORKSHEET P12
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "P12", "Pertemuan 12 · Build Tooling, npm, dan Framework SPA")
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(1)
    r_title = p_title.add_run("Dari Berkas Lepas ke Proyek Komponen yang Dibangun")
    r_title.font.name = "Carlito"
    r_title.font.size = Pt(16.0)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x37, 0x96)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("Pengembangan Aplikasi Berbasis Web · SIF302 · Semester Gasal 2026/2027")
    r_sub.font.name = "Carlito"
    r_sub.font.size = Pt(9.0)
    r_sub.font.color.rgb = RGBColor(0x5A, 0x60, 0x80)
    
    # Tabel Identitas
    t_id = doc.add_table(rows=2, cols=4)
    t_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_id.autofit = False
    set_table_borders(t_id, color="D0D7DE", sz="4")
    id_widths = [Inches(1.0), Inches(2.6), Inches(1.0), Inches(2.17)]
    id_data = [
        [("NAMA", "003796", True, "FFFFFF"), ("AMADEUS DHARMA AKBARINDRA", "FFFFFF", True, "1F2328"),
         ("NIM", "003796", True, "FFFFFF"), ("25523253", "FFFFFF", True, "1F2328")],
        [("KELAS", "003796", True, "FFFFFF"), ("SIF302 / B", "FFFFFF", False, "1F2328"),
         ("TANGGAL", "003796", True, "FFFFFF"), ("22 Oktober 2026", "FFFFFF", False, "1F2328")]
    ]
    for r_idx, row_cells_data in enumerate(id_data):
        for c_idx, (text, bg, is_bold, text_c) in enumerate(row_cells_data):
            cell = t_id.rows[r_idx].cells[c_idx]
            cell.width = id_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_bold and bg == "003796" else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Carlito"
            r.font.size = Pt(8.5)
            r.font.bold = is_bold
            if text_c == "FFFFFF":
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Dikerjakan sendiri, bukan berpasangan. Setiap mahasiswa memindahkan halaman profilnya sendiri ke bentuk komponen. Anda bebas memilih isi halaman; yang dinilai adalah apakah proyeknya berjalan, perilakunya tetap sama seperti Pertemuan 11, dan hasil buildnya bisa dibuka.",
        bold_prefix="",
        hex_bg="FEF2F2",
        border_color="991B1B"
    )
    
    # Capaian Pembelajaran & Yang Dinilai Hari Ini
    t_capaian = doc.add_table(rows=2, cols=2)
    style_standard_table(
        t_capaian,
        [Inches(2.7), Inches(4.07)],
        ["Capaian pembelajaran", "Yang dinilai hari ini"],
        [
            [
                "Mahasiswa mampu memindahkan halaman modular menjadi komponen, menjalankannya dengan npm dan Vite, lalu membangunnya untuk produksi.",
                "1. npm run dev membuka halaman, Console bersih.\n2. Token, CSS, dan HTML P04–P06 dipakai apa adanya.\n3. Empat komponen tersusun; satu di antaranya menerima data lewat props.\n4. npm run build menghasilkan folder dist/.\n5. npm run preview menjalankan hasil build itu."
            ]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "Aturan pengerjaan")
    t_aturan = doc.add_table(rows=7, cols=2)
    style_standard_table(
        t_aturan,
        [Inches(1.5), Inches(5.27)],
        ["Hal", "Ketetapan"],
        [
            ["Bentuk kerja", "Individu. Satu mahasiswa, satu proyek, satu topik."],
            ["Perkakas", "Node.js LTS dan npm (pekerjaan rumah dari Pertemuan 11). VS Code, terminal, dan peramban."],
            ["Berkas awal", "Seluruh berkas hasil Pertemuan 11: profil.html, enam berkas CSS, js/app.js, js/dom.js, js/data.js, js/state.js, js/events.js, dan js/api.js dari Pertemuan 10."],
            ["Tempat kerja", "Folder baru worksheet-p12/. Kerangka Vue dibuat di dalamnya, bukan di folder lain."],
            ["Waktu", "90 menit di kelas. Bila belum selesai, dilanjutkan di luar kelas paling lambat pukul 23.59 pada hari yang sama."],
            ["Penyerahan", "Seluruh isi worksheet-p12/ diunggah ke repositori GitHub Anda sendiri, di dalam folder semester ~/pabw/. Periksa halaman repositori Anda, bukan hanya folder di laptop."]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "Anggaran waktu 90 menit")
    t_anggaran1 = doc.add_table(rows=7, cols=3)
    style_standard_table(
        t_anggaran1,
        [Inches(0.9), Inches(3.87), Inches(2.0)],
        ["Menit", "Yang Anda kerjakan", "Kalau waktu habis, berhenti di sini"],
        [
            ["0–9 (A)", "Periksa node -v dan npm -v; perbaiki bila belum jalan; catat dua versinya", "Dua perintah berjalan, versi tercatat"],
            ["9–21 (B)", "Folder worksheet-p12/ kosong, npm create vite, npm install, tulis .gitignore", "npm run dev membuka halaman contoh Vue"],
            ["21–35 (C)", "Pindahkan aset P11 ke proyek Vue, impor CSS di main.js, hapus berkas contoh", "Halaman tampil sama seperti P11 lewat npm run dev"],
            ["35–52 (D)", "Pecah jadi empat komponen; satu komponen menerima props", "Empat komponen tersusun, perilaku sama"],
            ["52–68 (E)", "State reaktif dan localStorage lewat onMounted dan watch", "Pilihan terakhir tetap ada setelah muat ulang"],
            ["68–82 (F)", "npm run build, npm run preview, bandingkan ukuran, tiga kasus sulit", "dist/ terbentuk, situs jalan dari npm run preview"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    # ==========================================
    # HALAMAN 2: ANGGARAN WAKTU LANJUTAN
    # ==========================================
    t_anggaran2 = doc.add_table(rows=2, cols=3)
    style_standard_table(
        t_anggaran2,
        [Inches(0.9), Inches(3.87), Inches(2.0)],
        ["Menit", "Yang Anda kerjakan", "Kalau waktu habis, berhenti di sini"],
        [
            ["82–90 (G)", "Periksa satu per satu, satu baris untuk diingat, tiket keluar, penilaian mandiri, README", "Lembar G terisi dan tersimpan"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Selalu kerjakan berurutan. Lembar C mengandaikan lembar B sudah selesai, dan lembar E mengandaikan empat komponen sudah tersusun. Kolom paling kanan bukan batas waktu — hanya penunjuk laju.",
        bold_prefix="",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 3 & 4: LEMBAR A — MEMERIKSA PERKAKAS SEBELUM BEKERJA
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "A", "Lembar A — Memeriksa perkakas sebelum bekerja")
    
    make_callout_box(
        doc,
        "Seluruh pekerjaan hari ini berjalan di atas Node.js dan npm. Bila keduanya belum terpasang, tidak ada satu lembar pun yang bisa dikerjakan — jadi lembar ini dikerjakan paling awal, bukan paling akhir.",
        bold_prefix="▌ PERKAKAS DIPERIKSA SEBELUM BEKERJA",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "A.1 Dua perintah yang harus dijawab terminal")
    add_subtext(doc, "Ini pekerjaan rumah dari Pertemuan 11. Jalankan di terminal, di folder mana pun.")
    add_code_block(doc, "node -v\nnpm -v")
    
    t_a1 = doc.add_table(rows=3, cols=3)
    style_standard_table(
        t_a1,
        [Inches(1.5), Inches(1.8), Inches(3.47)],
        ["Perintah", "Versi yang muncul", "Catatan bila gagal"],
        [
            ["node -v", "v26.3.0", "Berhasil aktif normal (LTS siap pakai, melampaui syarat minimum 20.19+ / 22.12+)"],
            ["npm -v", "11.16.0", "Berhasil aktif normal (perintah dikenali secara global di terminal)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Versi Current bukan versi yang dipakai kelas ini. Yang LTS punya masa dukungan lebih panjang dan dipakai di komputer laboratorium. Perkakas yang dipakai hari ini (Vite) menuntut Node.js 20.19 ke atas atau 22.12 ke atas. Bila npm create vite di lembar B menolak dengan pesan tentang versi Node, naikkan Node.js LTS lebih dulu; jangan melanjutkan dengan versi lama.",
        bold_prefix="",
        hex_bg="FEF9C3",
        border_color="CA8A04"
    )
    
    add_heading_section(doc, "A.2 Periksa berkas Pertemuan 11 yang harus ada")
    add_subtext(doc, "Seluruh berkas di bawah ini dipindahkan ke proyek Vue di lembar C. Yang hilang akan terasa sebagai halaman yang kehilangan gaya atau data — bukan sebagai galat.")
    
    t_a2 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_a2,
        [Inches(2.5), Inches(1.2), Inches(3.07)],
        ["Berkas", "Ada?", "Jumlah baris"],
        [
            ["profil.html", "Ada [x]", "211 baris (kerangka lengkap dengan ID kontainer & form)"],
            ["Enam berkas CSS (P04–P06)", "Ada [x]", "563 baris (tokens, base, layout, komponen, tema, responsif)"],
            ["js/app.js, js/dom.js", "Ada [x]", "233 baris (app: 53 baris, dom: 180 baris)"],
            ["js/data.js, js/state.js, js/events.js", "Ada [x]", "328 baris (data: 102 baris, state: 112 baris, events: 114 baris)"],
            ["js/api.js (Pertemuan 10)", "Ada [x]", "73 baris (lapisan data fetch publik & JSON parsing)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "A.3 Periksa folder kerja")
    t_a3 = doc.add_table(rows=5, cols=3)
    style_standard_table(
        t_a3,
        [Inches(2.2), Inches(2.2), Inches(2.37)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Folder worksheet-p12/", "Ada di dalam ~/pabw, masih kosong", "Tersedia di dalam direktori repositori semester"],
            ["node_modules/ di worksheet-p11/", "Tidak ada — node_modules tidak pernah disalin", "Sesuai, node_modules tidak disalin sama sekali"],
            ["Berkas P11 sudah di-commit", "git status bersih sebelum mulai", "Sesuai, riwayat commit 5c9caf3 tersimpan bersih"],
            ["Nomor port npm run dev", "Dicatat; simpanan localStorage terikat port itu", "Port 5173 (localStorage pabw-p11 terikat pada port ini)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Sebelum menghapus apa pun, pastikan pekerjaan P11 sudah di-commit. Kerangka Vite yang dibuat di folder yang sudah berisi berkas akan menawarkan pilihan menghapus seluruh isi folder itu — dan yang tersisa hanya .git.",
        bold_prefix="",
        hex_bg="FEF2F2",
        border_color="991B1B"
    )
    
    t_sudah_a = doc.add_table(rows=6, cols=2)
    style_standard_table(
        t_sudah_a,
        [Inches(5.5), Inches(1.27)],
        ["Sudah dikerjakan", "Tanda"],
        [
            ["node -v dan npm -v keduanya menjawab dengan nomor versi", "[x] Selesai"],
            ["Node.js yang terpasang versi LTS, bukan Current", "[x] Selesai (v26.3.0)"],
            ["Seluruh berkas Pertemuan 11 lengkap dan sudah di-commit", "[x] Selesai"],
            ["Folder worksheet-p12/ sudah ada dan masih kosong", "[x] Selesai"],
            ["Dua nomor versi sudah saya catat, untuk ditulis di README di lembar G", "[x] Selesai"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_git_commit_box(doc, "A", "P12 — perkakas diperiksa, versi dicatat")
    add_check_verification(doc, "A", "buka terminal di ~/pabw dan jalankan node -v lalu npm -v. Bila keduanya menjawab nomor versi, pekerjaan hari ini bisa dimulai.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 5 & 6: LEMBAR B — MEMBUAT PROYEK BARU DI FOLDER BARU
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "B", "Lembar B — Membuat proyek baru di folder baru")
    
    make_callout_box(
        doc,
        "Kerangka Vue dibuat di folder yang masih kosong. Urutannya sengaja begitu: pada folder kosong tidak ada pertanyaan yang muncul dan tidak ada berkas yang bisa hilang.",
        bold_prefix="▌ PROYEK BARU, FOLDER BARU",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "B.1 Enam baris perintah, berurutan")
    add_subtext(doc, "Titik pada baris npm create vite@latest . berarti \"buat di folder ini\", bukan membuat subfolder baru. Perintahnya dijalankan setelah cd, bukan sebelumnya.")
    add_code_block(doc, "cd ~/pabw\nmkdir worksheet-p12\ncd worksheet-p12\nnpm create vite@latest . -- --template vue\nnpm install\nnpm run dev")
    
    add_heading_section(doc, "B.2 Bila pertanyaan pemasangan muncul")
    add_subtext(doc, "Pertanyaan itu hanya muncul bila foldernya sudah berisi berkas. Pada folder kosong perintahnya langsung menulis kerangka proyek. Bila tetap muncul, jawabannya sudah ditetapkan.")
    
    t_b2 = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_b2,
        [Inches(2.2), Inches(3.0), Inches(1.57)],
        ["Pilihan di layar", "Yang terjadi", "Yang Anda pilih"],
        [
            ["Cancel operation", "Perintah berhenti; tidak ada berkas yang disentuh", "Pilih ini"],
            ["Remove existing files and continue", "Seluruh isi folder dihapus, kecuali .git", "Jangan — pekerjaan P11 bisa hilang"],
            ["Ignore files and continue", "Folder tidak dikosongkan; nama yang bentrok ditimpa", "Jangan — berkas P11 bisa tertimpa"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Pada folder kosong baris pertama tabel itu tidak pernah ditawarkan. Inilah alasan lembar ini dikerjakan sebelum aset P11 dipindahkan — bukan sesudah.",
        bold_prefix="",
        hex_bg="DCFCE7",
        border_color="16A34A"
    )
    
    add_heading_section(doc, "B.3 Berkas dan folder yang terbentuk")
    add_subtext(doc, "Periksa isinya sebelum menulis satu baris pun kode Anda sendiri. Sesudah ini, proyek milik Anda.")
    
    t_b3 = doc.add_table(rows=8, cols=3)
    style_standard_table(
        t_b3,
        [Inches(1.8), Inches(2.5), Inches(2.47)],
        ["Berkas / folder", "Perannya", "Keterangan"],
        [
            ["index.html", "Halaman tunggal di akar proyek", "Nanti diisi dari profil.html"],
            ["src/main.js", "Titik masuk — tempat mengimpor CSS global", "Nanti ditambahi impor CSS P04–P06"],
            ["src/App.vue", "Komponen induk", "Nanti memuat empat komponen Anda"],
            ["src/components/", "Tempat komponen", "HelloWorld.vue di dalamnya akan dihapus"],
            ["public/", "Berkas statis yang dibaca lewat alamat", "Tempat berkas .json statis & poster dari P10"],
            ["vite.config.js", "Pengaturan Vite", "Dikonfigurasi preserveSymlinks untuk Windows"],
            ["package.json", "Daftar isi proyek", "Berisi scripts, dependencies, devDependencies"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "B.4 .gitignore sebelum commit pertama")
    add_subtext(doc, "Tulis ini sebelum commit pertama, bukan sesudah. Berkas yang sudah terlanjur masuk riwayat harus dihapus dengan git rm -r --cached, dan itu langkah yang tidak perlu.")
    add_code_block(doc, "# worksheet-p12/.gitignore\nnode_modules/\ndist/\n.env\n.DS_Store\n\n# JANGAN diabaikan:\n# package.json\n# package-lock.json")
    
    t_b4 = doc.add_table(rows=4, cols=2)
    style_standard_table(
        t_b4,
        [Inches(2.0), Inches(4.77)],
        ["Diabaikan", "Kenapa"],
        [
            ["node_modules/", "Isinya ribuan berkas dan bisa dipasang ulang dengan npm install"],
            ["dist/", "Hasil build, dibuat ulang dengan npm run build"],
            ["package-lock.json", "JUSTRU DISIMPAN — ia yang mengunci versi tepat seluruh paket"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "B.5 Yang saya lihat setelah npm run dev")
    t_b5 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_b5,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Halaman terbuka", "Halaman contoh Vue tampil di peramban", "Berhasil terbuka sempurna di browser"],
            ["Alamat di terminal", "Ada nomor port, misalnya localhost:5173 — dicatat", "http://localhost:5173/ (dicatat)"],
            ["Console", "Bersih, tanpa pesan merah", "Bersih, 0 pesan merah galat"],
            ["Folder node_modules/", "Ada di dalam worksheet-p12/", "Ada, memuat 34 paket terpasang"],
            ["Git status", "node_modules/ dan dist/ tidak terdaftar", "Sesuai, diabaikan sepenuhnya oleh .gitignore"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Hentikan server dengan Ctrl+C sebelum menutup terminal. Selama server berjalan, terminal itu terpakai dan perintah lain tidak bisa dijalankan di dalamnya.",
        bold_prefix="",
        hex_bg="FEF9C3",
        border_color="CA8A04"
    )
    
    add_git_commit_box(doc, "B", "P12 — kerangka Vue di folder kosong")
    add_check_verification(doc, "B", "hentikan server dengan Ctrl+C, lalu lihat halaman repositori GitHub Anda. node_modules/ tidak muncul di sana, dan package-lock.json justru muncul.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 7, 8, 9: LEMBAR C — MEMINDAHKAN ASET PERTEMUAN 11
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "C", "Lembar C — Memindahkan aset Pertemuan 11")
    
    make_callout_box(
        doc,
        "Lembar ini memindahkan pekerjaan Anda sendiri ke dalam proyek Vue. Tidak ada fitur baru, tidak ada perbaikan desain — hanya perpindahan tempat. Bila tampilan berubah, ada yang salah dalam pemindahan, bukan ada yang kurang.",
        bold_prefix="▌ PINDAH TANPA MENGUBAH PERILAKU",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "C.1 Apa yang berpindah, dan ke mana")
    t_c1 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_c1,
        [Inches(1.8), Inches(2.2), Inches(2.77)],
        ["Berkas P11", "Tujuannya di proyek Vue", "Yang diubah"],
        [
            ["profil.html", "index.html di akar proyek", "Baris <link> CSS dan <script> lama dibuang. Isinya hanya <div id=\"app\"></div> dan satu <script type=\"module\" src=\"/src/main.js\"></script>"],
            ["Enam berkas CSS", "src/ (tokens.css, base.css, dll.)", "Isinya tidak disentuh sama sekali, nilai token & kelas CSS 100% utuh"],
            ["js/app.js, js/dom.js", "Tidak ikut ke proyek baru", "render() dan sentuhan DOM digantikan <template> di lembar D. muat(), simpan(), aturUlang() dipindah ke src/simpanan.js dan App.vue"],
            ["js/data.js", "src/data.js", "Isinya tetap, setiap butir data proyek dipastikan memiliki ID unik"],
            ["js/state.js, js/events.js", "src/", "state.js dibungkus reactive() di lembar E; events.js menjadi @click dan @submit di lembar D"],
            ["js/api.js (Pertemuan 10)", "src/api.js", "Isinya tidak diubah; fetch tetap hanya di dalamnya"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "C.2 Enam impor CSS di main.js")
    add_subtext(doc, "Hanya di satu tempat ini berkas CSS diimpor. index.html tidak lagi memuat baris link ke berkas CSS mana pun. Keenam berkas diimpor berurutan seperti aslinya di P04 dan P06: tokens, base, layout, komponen, tema, lalu responsif.")
    
    add_code_block(doc, "// src/main.js\nimport { createApp } from 'vue'\nimport App from './App.vue'\n\n// gaya asli P04-P06, isinya tidak diubah dan urutannya tetap\nimport './tokens.css'\nimport './base.css'\nimport './layout.css'\nimport './komponen.css'\nimport './tema.css'\nimport './responsif.css'\n\ncreateApp(App).mount('#app')")
    add_code_block(doc, "<!-- index.html — hanya dua baris ini yang tinggal di dalam <body> -->\n<body class=\"page\">\n <div id=\"app\"></div>\n <script type=\"module\" src=\"/src/main.js\"></script>\n</body>")
    
    t_c2 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_c2,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Baris <link> di index.html", "Tidak ada lagi", "Bersih, tidak ada tag link CSS di index.html"],
            ["Enam berkas CSS", "Ada di src/ dan diimpor di main.js", "Lengkap 6 berkas diimpor berurutan di main.js"],
            ["Tampilan halaman", "Sama seperti P11 — warna, huruf, jarak", "Persis identik 100% tanpa perubahan gaya visual"],
            ["Console", "Bersih; tidak ada galat 404 berkas CSS", "Bersih, 0 galat 404 berkas CSS"],
            ["Berkas contoh Vue", "HelloWorld.vue dan isinya sudah dihapus", "Dihapus seluruhnya dari folder components/"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Gaya yang hilang hampir selalu berarti CSS-nya belum diimpor di main.js. Periksa Console: satu baris 404 sudah cukup untuk menjelaskan seluruh halaman yang tampak polos.",
        bold_prefix="",
        hex_bg="FEF9C3",
        border_color="CA8A04"
    )
    
    add_heading_section(doc, "C.3 Mengganti isi App.vue sebelum menjalankan npm run dev")
    add_subtext(doc, "Kerangka Vite masih membawa App.vue contoh yang memanggil HelloWorld.vue. Berkas HelloWorld.vue sudah Anda hapus, jadi digantikan sementara sebelum disusun penuh di Lembar D.")
    add_code_block(doc, "// src/App.vue — isi sementara\n<template>\n <h1>Halaman Anda</h1>\n <p>Bagian halaman dipindahkan di lembar D.</p>\n</template>")
    
    add_heading_section(doc, "C.4 Berkas js/api.js dari Pertemuan 10 ikut pindah")
    add_subtext(doc, "js/api.js ikut dipindahkan, bukan dibuang, dan isinya tidak diubah. Letaknya ditentukan bentuknya, bukan selera.")
    
    t_c4 = doc.add_table(rows=5, cols=2)
    style_standard_table(
        t_c4,
        [Inches(3.2), Inches(3.57)],
        ["Bentuk berkas", "Letaknya"],
        [
            ["Kode JavaScript yang diimpor", "src/ (src/api.js, src/data.js)"],
            ["Berkas .json statis", "public/ (public/film.json)"],
            ["Gambar yang dipakai di CSS", "public/ (public/images (3).jpg)"],
            ["Gambar yang diimpor kode", "src/assets/"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "C.5 Hasil yang diminta dari lembar ini")
    t_c5 = doc.add_table(rows=5, cols=3)
    style_standard_table(
        t_c5,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Halaman di npm run dev", "Terbuka tanpa galat; isinya sementara", "Terbuka normal, nol galat"],
            ["Berkas contoh Vite", "Sudah dihapus seluruhnya", "HelloWorld.vue & style.css contoh dihapus"],
            ["Data", "Ada di src/data.js dan dipakai halaman", "Ada di src/data.js dan terhubung"],
            ["Tampilan diubah?", "Tidak — token dan CSS P04–P06 apa adanya", "Tidak diubah, 100% gaya asli P11 dipertahankan"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_git_commit_box(doc, "C", "P12 — aset P11 dipindahkan apa adanya")
    add_check_verification(doc, "C", "buka halaman lewat npm run dev, lalu bandingkan dengan halaman P11 Anda. Warna, huruf, dan jaraknya harus sama.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 10, 11, 12: LEMBAR D — MEMECAH HALAMAN MENJADI EMPAT KOMPONEN
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "D", "Lembar D — Memecah halaman menjadi empat komponen")
    
    make_callout_box(
        doc,
        "Bagian halaman yang sudah ada di layar menjadi satu berkas. Aturannya sama dengan waktu memecah JavaScript di Pertemuan 11 — hanya bentuk berkasnya yang berbeda.",
        bold_prefix="▌ SATU KOMPONEN, SATU BAGIAN HALAMAN",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "D.0 Empat langkah, satu komponen setiap langkah")
    add_subtext(doc, "Langkah 1 — KepalaProfil.vue. Pindahkan bagian identitas dari profil.html, pasang di App.vue seperti D.3, lalu periksa: nama dan peran muncul, Console bersih.\nLangkah 2 — TombolFilter.vue. Pindahkan deretan tombol kategori. Periksa: tombol yang aktif berpindah saat diklik.\nLangkah 3 — DaftarProyek.vue. Pindahkan daftar kartu beserta penyaringannya. Periksa: kartu tampil, dan klik tombol kategori menyaring.\nLangkah 4 — FormKontak.vue. Pindahkan formulir kontak / tambah film. Periksa: huruf di kotak input tetap ada, dan kirim form mengosongkannya tanpa memuat ulang.")
    
    add_heading_section(doc, "D.1 Empat komponen dan tanggung jawabnya")
    t_d1 = doc.add_table(rows=5, cols=3)
    style_standard_table(
        t_d1,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Komponen", "Bagian halaman yang dipindahkan", "Data yang dipakainya"],
        [
            ["KepalaProfil.vue", "Bagian identitas di atas halaman & navbar", "props profil dari App.vue"],
            ["TombolFilter.vue", "Deretan tombol kategori & reset filter", "state bersama"],
            ["DaftarProyek.vue", "Daftar kartu proyek / film & pesan kosong", "state bersama"],
            ["FormKontak.vue", "Formulir tambah film / kontak", "ref lokal di komponennya (isi form)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "D.2 Padanan yang sudah Anda kenal")
    t_d2 = doc.add_table(rows=8, cols=2)
    style_standard_table(
        t_d2,
        [Inches(3.2), Inches(3.57)],
        ["Yang Anda tulis di P08–P11", "Bentuknya di Vue"],
        [
            ["daftarProyek.map(...) lalu ditempel", "v-for=\"proyek in state.daftarProyek\""],
            ["filter saat tombol diklik", "v-if atau daftar terhitung (computed)"],
            ["textContent = nilai", "{{ nilai }}"],
            ["setAttribute(\"src\", ...)", ":src=\"...\""],
            ["addEventListener(\"click\", ...)", "@click=\"...\""],
            ["Membaca input.value", "v-model=\"...\""],
            ["Wajib :key berisi id unik", "v-for=\"proyek in state.daftarProyek\" :key=\"proyek.id\" — id unik di src/data.js"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "D.3 Data dari induk ke anak lewat props")
    add_subtext(doc, "Hanya KepalaProfil yang menerima props, supaya tidak ada dua mekanisme yang harus dicerna sekaligus. Sisanya memakai state bersama.")
    
    add_code_block(doc, "// src/components/KepalaProfil.vue\n<script setup>\nconst props = defineProps({\n  profil: { type: Object, required: true },\n})\n</script>\n<template>\n  <h1>{{ props.profil.nama }}</h1>\n  <p>{{ props.profil.peran }}</p>\n</template>")
    add_code_block(doc, "// src/App.vue\nimport KepalaProfil from './components/KepalaProfil.vue'\nimport TombolFilter from './components/TombolFilter.vue'\nimport DaftarProyek from './components/DaftarProyek.vue'\nimport FormKontak from './components/FormKontak.vue'\n\nconst identitas = { nama: 'Amadeus Dharma Akbarindra', peran: 'Mahasiswa Informatika yang belajar front-end & mengulas film' }\n<KepalaProfil :profil=\"identitas\" />\n<TombolFilter /> <DaftarProyek /> <FormKontak />")
    
    t_d3 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_d3,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Empat komponen", "Satu berkas .vue untuk setiap bagian halaman", "Lengkap 4 berkas .vue di src/components/"],
            ["props KepalaProfil", "Terisi dari App.vue lewat :profil=\"identitas\"", "Terisi akurat via data identitas mahasiswa"],
            ["props diubah dari anak?", "Tidak — anak hanya membaca", "Sesuai, bersifat read-only murni"],
            ["Nama berkas dan impor", "Ditulis persis sama, huruf besar-kecilnya juga", "PascalCase konsisten persis di seluruh impor"],
            ["Console", "Bersih, tanpa peringatan props", "Bersih, tanpa ada warning atau error props"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Nama berkas yang berbeda hanya pada huruf besar-kecil akan berjalan di Windows dan gagal di Linux. Periksalah satu per satu selagi jumlahnya baru empat — jauh lebih lama bila sudah ratusan.",
        bold_prefix="",
        hex_bg="FEF2F2",
        border_color="991B1B"
    )
    
    add_heading_section(doc, "D.4 Pohon komponen yang saya susun")
    t_d4 = doc.add_table(rows=5, cols=3)
    style_standard_table(
        t_d4,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Komponen", "Di dalam komponen apa ia dipakai", "Menerima props?"],
        [
            ["KepalaProfil.vue", "App.vue", "Ya (menerima props profil)"],
            ["TombolFilter.vue", "App.vue", "Tidak (memakai state bersama)"],
            ["DaftarProyek.vue", "App.vue", "Tidak (memakai state bersama)"],
            ["FormKontak.vue", "App.vue", "Tidak (memakai ref lokal di komponennya)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_git_commit_box(doc, "D", "P12 — empat komponen tersusun")
    add_check_verification(doc, "D", "buka halaman lewat npm run dev dan periksa Console. Tidak ada peringatan tentang props, :key, atau nama komponen.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 13, 14, 15: LEMBAR E — MENYAMBUNGKAN KEMBALI KEADAAN DAN SIMPANANNYA
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "E", "Lembar E — Menyambungkan kembali keadaan dan simpanannya")
    
    make_callout_box(
        doc,
        "Ini bagian yang paling berbeda dari Pertemuan 11. Objek state Anda tetap dipakai, tetapi tidak ada lagi pemanggilan render(state) manual — Vue mengurusnya sendiri.",
        bold_prefix="▌ KEADAAN DI SATU TEMPAT, TAMPILAN MENGIKUTI SENDIRI",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "E.0 Tiga langkah, satu berkas setiap langkah")
    add_subtext(doc, "Langkah 1 — src/state.js. Bungkus objek state P11 dengan reactive() dan pindahkan pilihKategori() ke berkas yang sama.\nLangkah 2 — src/simpanan.js. Pindahkan simpan(), muat(), aturUlang(), KUNCI, dan VERSI dari js/app.js P11 apa adanya.\nLangkah 3 — App.vue. Pasang onMounted, watch, dan tombol atur ulang.")
    
    add_heading_section(doc, "E.1 render(state) menjadi state reaktif")
    t_e1 = doc.add_table(rows=6, cols=3)
    style_standard_table(
        t_e1,
        [Inches(2.0), Inches(2.2), Inches(2.57)],
        ["Pertemuan 11", "Pertemuan 12", "Catatan"],
        [
            ["const state = { ... }", "const state = reactive({ ... })", "reactive dipilih karena paling mirip objek state P11"],
            ["pilihKategori() memanggil render()", "pilihKategori() hanya mengubah state", "Vue menggambar ulang bagian yang berubah otomatis"],
            ["render(state) ditulis sendiri", "Tidak ada render() manual", "Blok template menggantikannya secara deklaratif"],
            ["Nilai tunggal disimpan di variabel", "const pesan = ref('')", "Dipakai hanya untuk nilai tunggal seperti isi form"],
            ["Hasil ambilNotisia() dari js/api.js", "Dimasukkan ke state di App.vue", "api.js tidak berubah sejak P10; fetch tetap di dalamnya"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_code_block(doc, "// src/state.js\nimport { reactive } from 'vue'\nimport { daftarProyek } from './data.js'\n\nexport const state = reactive({\n  daftarProyek: [...daftarProyek],\n  kategoriAktif: 'semua',\n})\n\nexport function pilihKategori(nilai) {\n  state.kategoriAktif = nilai // tanpa render() manual\n}\n\nexport function tambahFilm(filmBaru) {\n  state.daftarProyek.unshift(filmBaru)\n}")
    
    add_heading_section(doc, "E.2 localStorage P11 dipakai ulang lewat dua kait dan satu tombol")
    add_subtext(doc, "Berkas simpanan dari Pertemuan 11 (simpan, muat, aturUlang beserta KUNCI \"pabw-p11\" dan VERSI = 2) dipakai apa adanya di src/simpanan.js. onMounted membaca simpanan satu kali saat dipasang, watch menulis simpanan setiap kali state berubah, dan tombol atur ulang mengembalikan nilai bawaan.")
    
    add_code_block(doc, "// src/App.vue\nimport { onMounted, watch } from 'vue'\nimport { state } from './state.js'\nimport { muat, simpan, aturUlang } from './simpanan.js'\n\nonMounted(() => {\n  const isi = muat() // VERSI diperiksa di dalam muat(); bentuk lama dibuang\n  if (isi) state.kategoriAktif = isi.kategoriAktif\n})\n\nwatch(state, (nilai) => simpan(nilai), { deep: true })\n\nfunction aturUlangPilihan() {\n  aturUlang()\n  state.kategoriAktif = 'semua'\n  simpan(state)\n}")
    
    add_heading_section(doc, "E.3 Yang diperiksa")
    t_e3 = doc.add_table(rows=7, cols=3)
    style_standard_table(
        t_e3,
        [Inches(2.0), Inches(2.5), Inches(2.27)],
        ["Yang diperiksa", "Hasil yang benar", "Hasil yang saya dapat"],
        [
            ["Ganti kategori", "Hanya state yang berubah; tampilan mengikuti sendiri", "Kartu langsung tersaring instan tanpa render manual"],
            ["Tutup tab, lalu buka lagi", "Pilihan terakhir masih ada", "Pilihan kategori terakhir tetap tersimpan & aktif"],
            ["localStorage di Console", "Isinya berupa teks JSON, bukan [object Object]", "Teks JSON sah: {\"versi\":2,\"kategoriAktif\":\"Action\"}"],
            ["render(state) lama", "Sudah tidak ada di berkas mana pun", "Bersih, tidak ada lagi panggilan render manual"],
            ["Tombol atur ulang", "Kategori kembali ke \"semua\"; simpanan berisi nilai bawaan", "Kategori pulih ke \"semua\", simpanan direset bersih"],
            ["Console", "Bersih; tidak ada peringatan watch atau reaktivitas", "Bersih, 0 pesan peringatan reaktivitas"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "E.4 Bila simpanannya rusak")
    add_subtext(doc, "Kode muat() dari Pertemuan 11 tetap dipakai. Ujilah lagi di proyek baru ini, karena satu baris yang tertinggal saat pemindahan tidak terlihat sampai simpanannya rusak.")
    
    t_e4 = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_e4,
        [Inches(2.3), Inches(2.3), Inches(2.17)],
        ["Yang ditulis di Console", "Yang harus terjadi", "Hasil yang saya dapat"],
        [
            ["localStorage.setItem(\"pabw-p11\", \"bukan-json\")", "Halaman tetap terbuka; simpanan dibuang; nilai bawaan dipakai", "Tertangkap try/catch, simpanan dibuang, kembali ke default"],
            ["npm run dev dihentikan lalu dijalankan lagi", "Pilihan terakhir tetap ada", "Pilihan filter terakhir bertahan sempurna"],
            ["localStorage.clear()", "Halaman kembali ke nilai bawaan tanpa galat", "Pulih ke default tanpa menimbulkan pesan galat"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Bila simpanan tidak terbaca lagi setelah proyek dipindahkan, periksa kunci dan nomor port. Simpanan terikat pada alamat beserta portnya, dan Vite memilih port yang berbeda bila port sebelumnya sedang terpakai.",
        bold_prefix="",
        hex_bg="FEF9C3",
        border_color="CA8A04"
    )
    
    add_git_commit_box(doc, "E", "P12 — keadaan reaktif dan simpanan tersambung")
    add_check_verification(doc, "E", "ubah kategori, tutup tab, buka lagi. Pilihan terakhir Anda masih ada, dan Console bersih.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 16 & 17: LEMBAR F — MEMBANGUN DAN MENCOBA HASIL BUILD
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "F", "Lembar F — Membangun dan mencoba hasil build")
    
    make_callout_box(
        doc,
        "Sampai sekarang halaman Anda dijalankan oleh server pengembangan. Lembar ini menghasilkan versi yang benar-benar dikirim ke peramban — dengan berkas yang sudah digabung dan diperkecil.",
        bold_prefix="▌ SUMBER DAN HASIL DISIMPAN DI TEMPAT BERBEDA",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "F.1 Dua perintah yang berbeda gunanya")
    add_code_block(doc, "npm run build # menghasilkan folder dist/\nnpm run preview # menjalankan isi dist/")
    
    t_f1 = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_f1,
        [Inches(1.8), Inches(2.7), Inches(2.27)],
        ["Perintah", "Yang dijalankan", "Kapan dipakai"],
        [
            ["npm run dev", "Sumber Anda, apa adanya, dengan pemuatan ulang otomatis", "Selama proses mengerjakan / pengodean"],
            ["npm run build", "Menggabungkan dan memperkecil sumber menjadi dist/", "Sesudah pekerjaan selesai"],
            ["npm run preview", "Menjalankan isi dist/ seperti situs sungguhan", "Untuk memastikan hasil build benar sebelum rilis"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Hentikan npm run dev dengan Ctrl+C sebelum menjalankan npm run build. Keduanya memakai terminal yang sama, dan hasil build yang dibuat sambil server berjalan sering membingungkan karena alamatnya berbeda.",
        bold_prefix="",
        hex_bg="FEF9C3",
        border_color="CA8A04"
    )
    
    add_heading_section(doc, "F.2 Hasil build saya")
    t_f2 = doc.add_table(rows=7, cols=2)
    style_standard_table(
        t_f2,
        [Inches(3.2), Inches(3.57)],
        ["Yang diperiksa", "Hasil yang saya dapat"],
        [
            ["Folder dist/ terbentuk", "Terbentuk sempurna (exit code 0, waktu build ~2.99s - 7.7s)"],
            ["dist/index.html ada", "Ada (ukuran 0.44 kB / 442 byte)"],
            ["dist/assets/ berisi berkas .js dan .css", "Ada (index-DgvRWNVC.js dan index-6Nsuwklj.css)"],
            ["Nama berkas di dist/assets/ memuat huruf dan angka acak", "Ya, memuat hash acak otomatis untuk cache-busting"],
            ["Halaman jalan lewat npm run preview", "Jalan lancar di http://localhost:4173/"],
            ["Console bersih saat preview dibuka", "Bersih, seluruh komponen dan interaktivitas berjalan mulus"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "F.3 Ukuran sebelum dan sesudah build")
    add_subtext(doc, "Catat angkanya apa adanya. Yang dinilai adalah Anda mengenali hasilnya, bukan angkanya.")
    
    t_f3 = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_f3,
        [Inches(2.5), Inches(2.1), Inches(2.17)],
        ["Yang diukur", "Sebelum build", "Sesudah build"],
        [
            ["Ukuran seluruh berkas yang dikirim", "~450 kB (tersebar unbundled)", "83.8 kB mentah (hanya 32.2 kB gzip)"],
            ["Jumlah berkas yang dikirim peramban", "25+ berkas terpisah", "3 berkas utama (HTML, JS, CSS)"],
            ["Cara mengukurnya di DevTools", "Tab Network, muat ulang halaman", "Tab Network, muat ulang halaman"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "F.4 Tiga kasus sulit")
    add_subtext(doc, "Ketiganya muncul karena susunan folder, bukan karena Vue. Semuanya bisa Anda buat sendiri sekarang.")
    
    t_f4 = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_f4,
        [Inches(2.0), Inches(2.4), Inches(2.37)],
        ["Kasus", "Gejala", "Yang dilakukan"],
        [
            ["Nama berkas berbeda huruf besar-kecil", "Jalan di Windows, gagal di Linux atau CI", "Samakan huruf besar-kecil nama berkas dan nama impor, persis"],
            ["dist/ ter-commit", "Riwayat penuh berkas hasil build ber-hash", "git rm -r --cached dist/, lalu abaikan lewat .gitignore"],
            ["Gambar tidak muncul", "Gambar ada, tetapi alamatnya salah", "public/ dibaca lewat /nama.png; src/assets/ diimpor kode"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "F.5 Yang sengaja tidak saya commit")
    t_f5 = doc.add_table(rows=3, cols=2)
    style_standard_table(
        t_f5,
        [Inches(2.5), Inches(4.27)],
        ["Yang tidak di-commit", "Alasan saya"],
        [
            ["dist/", "Hasil build dibuat ulang kapan saja dengan npm run build"],
            ["node_modules/", "Isinya ribuan berkas dan bisa dipasang ulang dengan npm install"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "Simpan satu tangkapan layar halaman yang jalan dari npm run preview. Bukti itu yang Anda perlukan di lembar G, dan yang membuat pekerjaan Anda bisa diperiksa tanpa harus menjalankan terminal dosen.",
        bold_prefix="",
        hex_bg="DCFCE7",
        border_color="16A34A"
    )
    
    add_git_commit_box(doc, "F", "P12 — build produksi berhasil")
    add_check_verification(doc, "F", "buka alamat yang diberikan npm run preview. Halaman Anda jalan dari dist/, dan Console bersih.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 18, 19, 20: LEMBAR G — TIKET KELUAR DAN PENILAIAN MANDIRI
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "G", "Lembar G — Tiket keluar dan penilaian mandiri")
    
    make_callout_box(
        doc,
        "Dikerjakan sendiri, tanpa berdiskusi. Lembar ini mengukur pemahaman Anda, bukan kecepatan Anda menyalin jawaban teman.",
        bold_prefix="▌ TIKET KELUAR",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    qa_list = [
        ("1. Apa bedanya dependencies dan devDependencies? Sebutkan satu contoh untuk masing-masing dari proyek Anda.",
         "dependencies adalah pustaka yang dibutuhkan oleh aplikasi saat berjalan di lingkungan produksi (runtime) peramban pengguna; contohnya 'vue' (^3.5.13) yang menggerakkan reaktivitas dan komponen UI. Sedangkan devDependencies adalah perkakas bantu yang hanya dibutuhkan oleh pengembang selama tahap perancangan dan proses build (build-time); contohnya 'vite' (^6.0.7) dan '@vitejs/plugin-vue' (^5.2.1) yang mengompilasi Single File Component menjadi kode statis."),
        
        ("2. Mengapa package-lock.json wajib disimpan di repositori, sedangkan node_modules/ tidak?",
         "node_modules/ berisi puluhan ribu file fisik paket (ratusan MB) yang redundan karena dapat diunduh ulang kapan saja dengan 'npm install'. Sebaliknya, package-lock.json wajib disimpan karena mencatat manifes versi persis (exact version), pohon sub-dependensi, dan integrity hash seluruh paket. Hal ini menjamin bahwa siapapun atau mesin CI manapun yang menginstal paket akan mendapatkan struktur dependensi yang 100% identik tanpa konflik versi ('it works on my machine')."),
        
        ("3. Apa yang dilakukan perintah npm create vite@latest . -- --template vue? Apa arti titik pada perintah itu?",
         "Perintah ini mengunduh dan mengeksekusi alat pembangun resmi Vite terbaru untuk membuat proyek Single Page Application berbasis framework Vue 3. Tanda titik (.) berarti kerangka proyek dibuat di dalam 'direktori saat ini' (current working directory), bukan membuat subfolder baru. Tanda '--' digunakan oleh npm untuk meneruskan opsi '--template vue' langsung ke eksekutor Vite."),
        
        ("4. Di Pertemuan 11 Anda memanggil render(state) sendiri. Apa yang menggantikannya di Vue, dan siapa yang memanggilnya sekarang?",
         "Di Pertemuan 11, perenderan dilakukan secara imperatif manual lewat fungsi render(state). Di Vue, hal itu digantikan oleh blok <template> deklaratif yang terikat pada data reaktif (reactive() / ref()). Yang memanggil pembaruan tampilan sekarang adalah Sistem Reaktivitas bawaan Vue (Reactivity Engine & Virtual DOM), yang otomatis mendeteksi perubahan state dan memperbarui DOM yang relevan tanpa sentuhan manual pengembang."),
        
        ("5. Mengapa props tidak boleh diubah dari dalam komponen anak?",
         "Karena framework Vue memberlakukan prinsip One-Way Data Flow (Aliran Data Satu Arah). Data mengalir dari induk ke anak sebagai masukan read-only. Jika anak diizinkan mengubah props langsung, sumber kebenaran status (single source of truth) menjadi rusak, status komponen induk menjadi tidak sinkron, dan alur debugging menjadi tidak terduga (unpredictable side-effects). Jika anak ingin mengubah data, ia harus mengirim sinyal (emit event) ke induk."),
        
        ("6. Apa bedanya npm run dev dan npm run preview? Kapan Anda memakai masing-masing?",
         "npm run dev menjalankan server pengembangan lokal (Vite Dev Server) yang menyajikan kode sumber asli secara instan dengan Hot Module Replacement (HMR); dipakai sehari-hari SELAMA masa pengodean. npm run preview menjalankan server lokal statis untuk menyajikan dan menguji berkas ter-minifikasi hasil build di folder dist/; dipakai SESUDAH build selesai guna memvalidasi perilaku nyata aplikasi sebelum di-deploy ke server produksi."),
        
        ("7. Mengapa dist/ diabaikan sedangkan hasilnya justru yang dijalankan peramban di produksi?",
         "Folder dist/ adalah berkas artefak turunan hasil build otomatis (generated build artifact), bukan kode sumber asli. Menyimpan dist/ di Git akan mengotori riwayat komit dengan berkas ter-minifikasi ber-hash acak yang membengkakkan repositori. Di industri modern (CI/CD), server produksi akan menarik kode sumber dari Git lalu menjalankan 'npm run build' sendiri secara otomatis di lingkungan deployment.")
    ]
    
    for q, a in qa_list:
        add_heading_section(doc, q)
        t_box = doc.add_table(rows=1, cols=1)
        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_box.autofit = False
        c = t_box.rows[0].cells[0]
        c.width = Inches(6.77)
        set_cell_background(c, "F8FAFC")
        set_cell_margins(c, top=50, bottom=50, left=70, right=70)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(a)
        r.font.name = "Carlito"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        
    add_heading_section(doc, "G.1 Periksa satu per satu")
    t_g1 = doc.add_table(rows=11, cols=2)
    style_standard_table(
        t_g1,
        [Inches(5.5), Inches(1.27)],
        ["Kriteria Pemeriksaan", "Status"],
        [
            ["Perkakas node -v dan npm -v menjawab; versinya tercatat di README", "[x] Terpenuhi"],
            ["Proyek worksheet-p12/ dibuat dari folder terstruktur; npm run dev membuka halaman tanpa galat", "[x] Terpenuhi"],
            ["Pemindahan halaman tampil sama seperti P11; token dan CSS P04–P06 tidak diubah", "[x] Terpenuhi"],
            ["Komponen empat komponen tersusun; KepalaProfil menerima props profil; nama berkas dan impor ditulis sama", "[x] Terpenuhi"],
            ["Keadaan state memakai reactive(); tidak ada render(state) manual lagi", "[x] Terpenuhi"],
            ["Simpanan onMounted memuat dan watch menyimpan; pilihan terakhir bertahan setelah tab ditutup; tombol atur ulang mengembalikan nilai bawaan", "[x] Terpenuhi"],
            ["Build npm run build menghasilkan dist/; npm run preview menjalankannya; dua ukuran sudah dibandingkan", "[x] Terpenuhi"],
            ["Git .gitignore memuat node_modules/ dan dist/; package-lock.json ikut di-commit; lebih dari satu commit", "[x] Terpenuhi"],
            ["Deklarasi AI README menyebutkan bagian mana yang dibantu AI dan bagian mana yang saya kerjakan sendiri", "[x] Terpenuhi"],
            ["Lapisan data js/api.js dari Pertemuan 10 ikut pindah, isinya tidak diubah, dan fetch hanya ada di dalamnya", "[x] Terpenuhi"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "G.2 Satu baris untuk diingat")
    add_subtext(doc, "Tulis satu kalimat — bukan untuk dosen, untuk Anda sendiri saat mengerjakan proyek berikutnya — tentang hal yang hari ini paling sering membuat Anda salah.")
    make_callout_box(
        doc,
        "\"Di era framework komponen, kita tidak lagi mengutak-atik DOM secara manual: cukup ubah datanya secara reaktif, dan biarkan framework yang memperbarui tampilannya secara presisi.\"",
        bold_prefix="",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    add_heading_section(doc, "G.3 Penilaian mandiri")
    add_subtext(doc, "Isi perkiraan nilai Anda apa adanya sebelum menyerahkan lembar ini. Kolom Bukti diisi nama berkas atau baris kode yang bisa dibuka dosen.\n[x] Pengungkapan AI README menyebutkan bagian mana yang dibantu AI dan bagian mana yang saya kerjakan sendiri\n[x] Keaslian topik dan isi halaman ini milik saya sendiri, tidak menyalin pekerjaan mahasiswa lain")
    
    t_g3 = doc.add_table(rows=7, cols=4)
    style_standard_table(
        t_g3,
        [Inches(2.5), Inches(0.8), Inches(0.8), Inches(2.67)],
        ["Bagian", "Bobot", "Nilai", "Bukti Berkas / Baris Kode"],
        [
            ["Proyek Vue berjalan", "20%", "20", "package.json, vite.config.js, console bersih"],
            ["Pemindahan tanpa mengubah perilaku", "25%", "25", "src/*.css, src/api.js, src/data.js apa adanya"],
            ["Pemecahan komponen dan props", "20%", "20", "src/components/*.vue (KepalaProfil :profil)"],
            ["Build produksi dan preview", "25%", "25", "dist/ terbentuk sempurna, preview port 4173"],
            ["Kebersihan src/, .gitignore, deklarasi AI", "10%", "10", ".gitignore benar, package-lock di-commit, README"],
            ["TOTAL", "100%", "100", "Sangat Baik (Lengkap & Sempurna)"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_heading_section(doc, "G.4 Catatan untuk pengampu")
    add_subtext(doc, "Bagian tersulit hari ini:\nMenjaga sinkronisasi persistensi localStorage menggunakan lifecycle hook onMounted dan deep watch tanpa memicu rendering berulang, serta penanganan path Windows yang memuat karakter khusus.")
    add_subtext(doc, "Yang ingin saya bahas di kelas berikutnya:\nManajemen state berskala kompleks menggunakan Pinia serta navigasi multi-halaman berbasis Vue Router.")
    add_subtext(doc, "Kendala perkakas yang saya alami (versi Node.js, port, atau lainnya):\nKarakter hash (#) pada path direktori sempat dicegat Vite dev server, berhasil diselesaikan secara mandiri lewat virtual drive subst V: dan konfigurasi preserveSymlinks: true.")
    
    # Tanda tangan & Box CPMK
    p_ttd = doc.add_paragraph()
    p_ttd.paragraph_format.space_before = Pt(4)
    p_ttd.paragraph_format.space_after = Pt(4)
    r_ttd = p_ttd.add_run("Tanda tangan: Amadeus Dharma Akbarindra          Tanggal: 22 Oktober 2026          Jam selesai: 16.56 WIB")
    r_ttd.font.name = "Carlito"
    r_ttd.font.size = Pt(8.5)
    r_ttd.font.bold = True
    
    make_callout_box(
        doc,
        "Pertemuan 13 adalah Ujian CPMK Blok 2. Bahan ujiannya DOM, event, dan API JSON — seluruhnya vanilla. Framework tidak diujikan, jadi simpan proyek ini rapi tanpa perlu menghafal Vue.",
        bold_prefix="",
        hex_bg="FEF2F2",
        border_color="991B1B"
    )
    
    add_git_commit_box(doc, "G", "P12 — periksa, lalu simpan")
    add_check_verification(doc, "G", "buka README di repositori Anda. Bagian Pertemuan 12 memuat deklarasi AI, dua nomor versi Node.js, dan dua ukuran sebelum serta sesudah build.")
    
    doc.add_page_break()
    
    # ==========================================
    # HALAMAN 21, 22, 23: LAMPIRAN A, B, DAN GIT
    # ==========================================
    make_worksheet_header(doc, "WORKSHEET", "H", "Lampiran A — Rubrik penilaian")
    
    add_heading_section(doc, "RUBRIK PENILAIAN (total 100)")
    t_rubrik = doc.add_table(rows=6, cols=4)
    style_standard_table(
        t_rubrik,
        [Inches(1.8), Inches(0.6), Inches(2.2), Inches(2.17)],
        ["Aspek", "Bobot", "Nilai 100 — sangat baik", "Nilai 65 — memadai"],
        [
            ["Proyek Vue berjalan", "20%", "Proyek dibuat terstruktur; npm run dev membuka halaman; Console bersih; nol galat impor", "Proyek berjalan, tetapi Console memuat peringatan yang tidak ditangani"],
            ["Pemindahan tanpa mengubah perilaku", "25%", "Token, CSS, dan HTML P04–P06 dipakai apa adanya; tampilan sama seperti P11; js/api.js dari P10 ikut dipindahkan dan dipanggil", "Halaman jalan, tetapi ada nilai token yang ditulis ulang atau bagian hilang"],
            ["Pemecahan komponen dan props", "20%", "Empat komponen tersusun dengan satu tanggung jawab; KepalaProfil menerima props profil; props dibaca saja; penamaan konsisten", "Komponen terpisah, tetapi props diubah dari anak atau satu komponen memikul dua tugas"],
            ["Build produksi dan preview", "25%", "npm run build menghasilkan dist/; npm run preview menjalankannya; dua ukuran dibandingkan dan dicatat", "dist/ terbentuk, tetapi belum pernah dibuka lewat npm run preview atau ukuran belum dicatat"],
            ["Kebersihan, .gitignore, dan bukti", "10%", "src/ rapi tanpa berkas contoh; node_modules/ dan dist/ diabaikan; lockfile di-commit; deklarasi AI terisi; lebih dari satu commit", "Repo bersih tetapi deklarasi AI belum lengkap, atau hanya satu commit"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_subtext(doc, "Yang dinilai adalah proyek Anda sendiri. Isi halaman boleh berbeda antar mahasiswa; yang diperiksa adalah apakah setiap bagian bekerja sebagaimana dimaksud dan apakah pemindahannya tidak mengubah perilaku halaman P11.", italic=True)
    
    add_heading_section(doc, "Lampiran B — Rujukan cepat dan kalau macet")
    make_callout_box(
        doc,
        "REFERENSI CEPAT NPM, VITE, DAN KOMPONEN hafalkan yang ini",
        bold_prefix="▌",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    t_rujukan = doc.add_table(rows=13, cols=3)
    style_standard_table(
        t_rujukan,
        [Inches(2.0), Inches(2.7), Inches(2.07)],
        ["Yang dipakai", "Perannya", "Kapan dipakai"],
        [
            ["npm create vite@latest . -- --template vue", "Membuat kerangka proyek di folder ini", "Awal pembuatan proyek"],
            ["npm install", "Memasang paket sesuai package.json", "Setelah kerangka dibuat atau repo di-clone"],
            ["npm run dev", "Menjalankan server pengembangan", "Selama masa pengerjaan / pengodean"],
            ["npm run build", "Menghasilkan folder dist/ produksi", "Sesudah pekerjaan selesai"],
            ["npm run preview", "Menjalankan isi dist/", "Memastikan hasil build benar"],
            ["package.json", "Daftar isi proyek", "Mencatat scripts dan dependensi"],
            ["package-lock.json", "Mengunci versi tepat seluruh paket", "Disimpan di Git untuk konsistensi tim"],
            ["reactive({ })", "Objek state yang tampilan mengikutinya", "State bersama banyak properti"],
            ["ref(nilai)", "Nilai tunggal yang reaktif", "Nilai tunggal seperti kolom form"],
            ["defineProps({ })", "Menyatakan data yang diterima dari induk", "Komponen anak menerima data"],
            ["v-for / v-if", "Padanan map dan filter deklaratif", "Perenderan daftar dan kondisi"],
            ["v-model", "Mengikat isi form dua arah", "Sinkronisasi input form dengan ref"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    make_callout_box(
        doc,
        "KALAU MACET cari gejalanya di sini",
        bold_prefix="▌",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    t_macet = doc.add_table(rows=11, cols=3)
    style_standard_table(
        t_macet,
        [Inches(2.0), Inches(2.4), Inches(2.37)],
        ["Gejala", "Penyebab yang paling sering", "Yang dilakukan"],
        [
            ["Current directory is not empty", "Perintah dijalankan di folder yang sudah berisi berkas", "Pilih Cancel, lalu buat folder baru yang kosong"],
            ["node -v tidak dikenali", "PATH belum memuat Node.js, atau terminal lama", "Tutup dan buka ulang terminal; pasang ulang bila tetap gagal"],
            ["Halaman polos tanpa warna", "CSS belum diimpor di main.js", "Tambahkan import './tokens.css' dan seterusnya di main.js"],
            ["Console: 404 berkas CSS", "Jalur impor CSS salah", "Periksa nama berkas di src/ dan samakan"],
            ["Failed to resolve import", "Nama berkas berbeda huruf besar-kecil", "Samakan nama berkas, nama impor, dan pemakaiannya"],
            ["props tidak muncul isinya", "props diakses tanpa props. atau didestrukturisasi", "Akses lewat props.nama atau toRefs"],
            ["Daftar tampil tetapi aneh saat berubah", ":key belum dipasang", "Tambahkan :key berisi id unik"],
            ["Simpanan tidak terbaca lagi", "Port server berubah, atau kunci localStorage berbeda", "Pakai satu kunci tetap ('pabw-p11') dan catat nomor portnya"],
            ["dist/ ikut ter-commit", "dist/ belum ada di .gitignore", "git rm -r --cached dist/, lalu tambahkan ke .gitignore"],
            ["Gambar tidak muncul", "Gambar di src/assets/ tapi dipanggil lewat alamat", "Impor gambarnya di kode atau pindahkan ke public/"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    
    add_subtext(doc, "Pertemuan 13 adalah Ujian CPMK Blok 2 — bahan ujiannya DOM, event, dan API JSON, seluruhnya vanilla. Pertemuan 14 melanjutkan dari folder dist/ yang Anda hasilkan hari ini, jadi simpan pekerjaan ini rapi di repositori Anda.", italic=True)
    
    make_callout_box(
        doc,
        "LAMPIRAN — MENYIMPAN PEKERJAAN ANDA KE GIT sekali saja, bukan tiap pertemuan\nPanduan Git dan GitHub yang sudah Anda pakai sejak Pertemuan 3 masih berlaku seluruhnya; bagian ini hanya mengulang tiga perintah yang paling sering dipakai, karena sekarang ada folder yang tidak boleh ikut disimpan.",
        bold_prefix="▌",
        hex_bg="EEF2FA",
        border_color="003796"
    )
    
    t_git_app = doc.add_table(rows=4, cols=3)
    style_standard_table(
        t_git_app,
        [Inches(1.8), Inches(3.0), Inches(1.97)],
        ["Perintah", "Artinya", "Kapan dijalankan"],
        [
            ["git add .", "Menandai seluruh perubahan di folder ini, kecuali yang ada di .gitignore", "Setiap kali selesai satu lembar"],
            ["git commit -m \"pesan\"", "Menyimpan perubahan itu sebagai satu riwayat", "Langsung setelah git add"],
            ["git status", "Menampilkan apa yang akan disimpan", "Sebelum commit pertama"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    add_subtext(doc, "Titik pada git add . Titik itu berarti \"seluruh berkas di folder ini\". Berkat .gitignore, node_modules/ dan dist/ tidak ikut meskipun titik itu dipakai.")
    
    add_heading_section(doc, "Berapa kali harus menyimpan?")
    add_subtext(doc, "Tidak ada angka yang benar, tetapi patokannya begini: bila Anda sudah menyelesaikan satu hal yang berarti dan bisa Anda jelaskan, simpanlah. Pada pertemuan ini minimal tiga kali: setelah proyek berjalan, setelah empat komponen tersusun, dan setelah build berhasil.")
    
    t_git_case = doc.add_table(rows=4, cols=2)
    style_standard_table(
        t_git_case,
        [Inches(3.2), Inches(3.57)],
        ["Kalau ini terjadi", "Yang Anda lakukan"],
        [
            ["Sudah menulis banyak, belum pernah commit", "Simpan sekarang juga"],
            ["node_modules/ terlanjur ter-commit", "git rm -r --cached node_modules/, commit, lalu abaikan"],
            ["Salah commit, belum push", "Perbaiki lalu commit lagi — jangan hapus repositori"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    make_callout_box(
        doc,
        "Yang dinilai bukan jumlah commit. Menyimpan sering membuat pekerjaan Anda tidak bergantung pada satu berkas di satu komputer — dan pekerjaan yang tidak ada di repositori tidak dapat dinilai.",
        bold_prefix="",
        hex_bg="DCFCE7",
        border_color="16A34A"
    )
    
    doc.save(output_path)
    print(f"Dokumen Word berhasil dibuat dan disimpan ke: {output_path}")

if __name__ == "__main__":
    out_file = "AMADEUS DHARMA AKBARINDRA - 25523253 - 12. Worksheet_PABW_Pertemuan12 - ID.docx"
    generate_p12_docx(out_file)
